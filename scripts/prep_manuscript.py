import argparse
import base64
import bs4
import datetime
import docx
import io
import docx.oxml
import docx.document
import mammoth
import pathlib
from PIL import Image
import sys
import os
import typing as t


# Get the parent directory and add it to sys.path
current_dir = os.path.dirname(os.path.abspath(__file__))
parent_dir = os.path.dirname(current_dir)
sys.path.append(parent_dir)

from src._common.constants import (
    MAMMOTH_STYLE_MAP,
    METADATA_FIELDS,
    FIELD_DIGITAL_SOURCE,
    FIELD_PAPER_SOURCE,
)


os.environ["PYTHONHASHSEED"] = "0"
parser = argparse.ArgumentParser(description="...")
parser.add_argument("book_path")


digital_input_field = next(
    f for f in METADATA_FIELDS if f["key"] == FIELD_DIGITAL_SOURCE
)
paper_input_field = next(f for f in METADATA_FIELDS if f["key"] == FIELD_PAPER_SOURCE)
iscolor_field = next(f for f in METADATA_FIELDS if f["key"] == "iscolor")


def _main() -> None:
    """
    This script does the following:
    1. Extract images from docx files to an assets folder
    2. Prepares Metadata Tables from docx files and partially filles them
    """
    args = parser.parse_args()
    book_path_arg = args.book_path.split("=")
    assert len(book_path_arg) == 2
    book_path = pathlib.Path(book_path_arg[1])
    metadata_tables: dict[str, list[dict[str, str]]] = {}
    assert book_path.exists(), f"{book_path} does not exist"
    (book_path / "assets").mkdir(exist_ok=True)
    for fp in book_path.glob("*"):
        if not fp.exists() or fp.is_dir():
            continue
        elif fp.name.startswith("~$"):
            continue
        elif fp.suffix == ".docx":
            metadata_tables[fp.name] = []
            with open(fp, "rb") as f:
                html_str = mammoth.convert_to_html(f, style_map=MAMMOTH_STYLE_MAP).value
                html_soup = bs4.BeautifulSoup(html_str, "html.parser")
                imgs = html_soup.find_all("img")
                if imgs:
                    target_dir = book_path / "assets" / fp.stem
                    target_dir.mkdir(exist_ok=True)
                for i, img in enumerate(imgs, start=1):
                    img_metadata = _extract_metadata_from_img(
                        *_extract_img_from_docx(i, target_dir, img)
                    )
                    metadata_tables[fp.name].append(img_metadata)

    notes_fn = (
        book_path
        / "assets"
        / f"{datetime.date.today().isoformat()}_editorial_notes.docx"
    )
    _write_editorial_notes_to_docx(notes_fn, metadata_tables)
    print("end")


def _extract_img_from_docx(
    i: int,
    target_dir: pathlib.PurePath,
    img: bs4.element.Tag,
) -> t.Tuple[str, str, str]:
    src: str = img["src"]
    alttext: str = img.get("alt", "")
    start = src.index("/") + 1
    end = src.index(";", start)
    suffix = src[start:end]
    encodded_img = src.split(",")[-1]
    fn = "{i:03}.{suffix}".format(i=i, suffix=suffix)
    img_fn: str = "/".join((target_dir / fn).parts[-3:])
    pil_img = Image.open(io.BytesIO(base64.b64decode(encodded_img)))
    try:
        pil_img.save(target_dir / fn, quality=100, subsampling=0)
        # print(f"saved {(target_dir / fn)}")
    except (KeyError, ValueError) as e:
        print(f"Error saving image {(target_dir / fn)}: {e}")
        # new_fp = (target_dir / fn).with_suffix(".png")
        # img_fn: str = "/".join(new_fp.parts[-3:])
        # pil_img.save(new_fp, quality=100, subsampling=0)
    return img_fn, encodded_img, alttext


def _extract_metadata_from_img(
    img_fn: str, encodded_img: str, alttext: str
) -> dict[str, str]:
    img_metadata = {}
    for field in METADATA_FIELDS:
        img_metadata["img"] = encodded_img
        if field["is_required"] and field["key"] == "id":
            img_metadata[field["show"]] = str(hash(encodded_img))
        elif field["is_required"] and field["key"] == "alttext":
            img_metadata[field["show"]] = alttext
        elif field["is_required"] and field["key"] in (
            FIELD_DIGITAL_SOURCE,
            FIELD_PAPER_SOURCE,
        ):
            img_metadata[digital_input_field["show"]] = img_fn
            img_metadata[paper_input_field["show"]] = img_fn
        elif field["is_required"]:
            img_metadata[field["show"]] = ""
    img_metadata[iscolor_field["show"]] = ""
    return img_metadata


def _write_editorial_notes_to_docx(
    fn: t.Union[str, pathlib.PurePath],
    metadata_tables: dict[str, list[dict[str, str]]],
) -> None:
    doc = docx.Document()
    _write_metadata_tables_to_docx(doc, metadata_tables)
    doc.save(fn)


def add_hyperlink(
    paragraph: docx.text.paragraph.Paragraph, text: str, url: str
) -> docx.oxml.text.hyperlink.CT_Hyperlink:
    """
    Source: https://stackoverflow.com/a/47666747/8007403
    """
    # This gets access to the document.xml.rels file and gets a new relation id value
    part = paragraph.part
    r_id = part.relate_to(
        url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    # Create the w:hyperlink tag and add needed values
    hyperlink = docx.oxml.shared.OxmlElement("w:hyperlink")
    hyperlink.set(
        docx.oxml.shared.qn("r:id"),
        r_id,
    )
    # Create a new run object (a wrapper over a 'w:r' element)
    new_run = docx.text.run.Run(docx.oxml.shared.OxmlElement("w:r"), paragraph)
    new_run.text = text
    # Set the run's style to the builtin hyperlink style, defining it if necessary
    # new_run.style = get_or_create_hyperlink_style(part.document)
    # Alternatively, set the run's formatting explicitly
    new_run.font.color.rgb = docx.shared.RGBColor(0, 0, 255)
    new_run.font.underline = True
    # Join all the xml elements together
    hyperlink.append(new_run._element)
    paragraph._p.append(hyperlink)
    return hyperlink


def _write_metadata_tables_to_docx(
    doc: docx.document.Document, metadata_tables: dict[str, dict[str, str]]
) -> None:
    doc.add_heading("Figures metadata", level=1)
    metadata_descs = {}
    for field in METADATA_FIELDS:
        metadata_descs[field["show"]] = field["desc"]
    for docx_fp, tables_data in metadata_tables.items():
        if not tables_data:
            continue
        doc.add_heading(docx_fp, level=2)
        for table_data in tables_data:
            table_data: dict
            _img = table_data.pop("img")
            try:
                doc.add_picture(
                    io.BytesIO(base64.b64decode(_img)), width=docx.shared.Inches(2)
                )
            except Exception as e:
                doc.add_paragraph("Error")
            tbl = doc.add_table(len(table_data), cols=2, style="Table Grid")
            for i, (key, value) in enumerate(table_data.items()):
                tbl.cell(i, 0).text = key
                tbl.cell(i, 1).text = value

            doc.add_paragraph("")
            doc.add_paragraph("Meaning of the fields:")
            for field_name in table_data.keys():
                doc.add_paragraph(
                    field_name + " - " + metadata_descs[field_name], style="List Bullet"
                )
            doc.add_paragraph("")
            p = doc.add_paragraph(
                "Make sure that the figure is associated with a caption. The caption must be structured in a way described by "
            )
            add_hyperlink(
                p,
                "the guidelines.",
                "https://cpy-dgtmon-dev-staticws-wip.s3.eu-west-1.amazonaws.com/read/guidelines-for-drafting-manuscripts/02_introduction_to_docx.html#how-to-structure-image-captions",
            )
            doc.add_page_break()


if __name__ == "__main__":
    _main()
