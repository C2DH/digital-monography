import argparse
import bs4
import docx
import pathlib


parser = argparse.ArgumentParser(description="...")
parser.add_argument("book_path", type=str)

"""
python scripts/prepare_and_build.py book_path="/mnt/c/Users/Pawel.kaminski/OneDrive - University of Luxembourg/Books/guidelines-for-drafting-manuscripts"/
python scripts/out_paper.py  book_path=/tmp/efs/guidelines-for-drafting-manuscripts/_html/
"""

EXLUDED_TAGS = (
    "script",
    # for now all known instances are about custom directives, e.g.
    # <span hidden="hidden">% :::{dgt-mon-asset} ed9c2212-4ff1-4ae5-a77c-cd949abf7d91 :::</span>
    "span",
)


def is_div(elem: bs4.element.Tag) -> bool:
    return elem.name == "div"


def is_h(elem: bs4.element.Tag) -> bool:
    return elem.name.startswith("h") and elem.name[1:].isdigit()


def is_p(elem: bs4.element.Tag) -> bool:
    return elem.name == "p"


def div_handler(doc: docx.document.Document, elem: bs4.element.Tag) -> None:
    # NOTE: This has to be handled.
    # For now all known instances include bibliographies
    # <div><h2>Bibliography</h2><ol><li><small>Puckett, J., 2017. <i>Zotero: A Guide For Librarians, Researchers, and Educators</i> 2nd edition., Assoc of College &amp; Research Libraries.</small></li><li><small>Fickers, A. &amp; Tatarinov, J. eds., 2022. <i>Digital History and Hermeneutics</i> first edition. A. Fickers &amp; J. Tatarinov, eds., De Gruyter Oldenbourg. Available at: https://doi.org/10.1515/9783110723991.</small></li></ol></div>
    return None


def h_handler(doc: docx.document.Document, elem: bs4.element.Tag) -> None:
    doc.add_heading(elem.text, level=int(elem.name[1:]))


def p_handler(doc: docx.document.Document, elem: bs4.element.Tag) -> None:
    doc.add_paragraph(elem.text)


checks_and_handlers = [
    (is_div, div_handler),
    (is_h, h_handler),
    # (is_ol, ol_handler), # ol/li
    (is_p, p_handler),
    # (is_blockquote, blockquote_handler),
    # (is_pre, pre_handler),
    # (is_table, table_handler),
    # (is_ul, ul_handler), # ul/li
]


if __name__ == "__main__":
    args = parser.parse_args()
    book_path_arg = args.book_path.split("=")
    assert len(book_path_arg) == 2
    book_path = pathlib.Path(book_path_arg[1])
    # toprint = set()
    for src in book_path.glob("*"):
        print("src ", src)
        if not src.exists() or src.is_dir() or src.name.startswith("~$"):
            continue
        elif src.suffix == ".html":
            dst = src.with_suffix(".docx")
            with src.open("r") as f:
                html_str = f.read()
                html_soup = bs4.BeautifulSoup(html_str, "html.parser")
                doc = docx.Document()
                for elem in html_soup.children:
                    if not isinstance(elem, bs4.element.Tag):
                        continue
                    elif elem.name in EXLUDED_TAGS:
                        continue
                    for check, handler in checks_and_handlers:
                        if check(elem):
                            handler(doc, elem)
                    # elif elem.name.startswith("h") and elem.name[1:].isdigit():
                    #     doc.add_heading(elem.text, level=int(elem.name[1:]))
                    # elif elem.name == "p":
                    #     doc.add_paragraph(elem.text)
                    # else:
                    #     toprint.add(elem.name)
                doc.save(f"/tmp/efs/{dst.name}")
        # print(toprint)
